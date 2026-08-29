"""
Service for ingesting and exporting alternative Raw Data formats (Weekly Matrix layout).
Matches the layout, structure, and headers of Raw-data-format.XLS.
"""
import os
import re
from datetime import datetime, timedelta, date
from typing import Tuple, Optional, Dict, Any, List
import pandas as pd
import openpyxl

FRENCH_DAYS = ['lundi', 'mardi', 'mercredi', 'jeudi', 'vendredi', 'samedi', 'dimanche']

class RawDataService:
    """Ingress and Egress layer for alternative Weekly Matrix Raw Data format"""

    @staticmethod
    def parse_raw_data_file(path: str) -> Tuple[Optional[pd.DataFrame], Optional[Dict[str, Any]]]:
        """
        Parse an alternative raw data file (.xls or .xlsx in weekly matrix format).
        Reconstructs the standard application DataFrame and metadata.
        """
        try:
            if not os.path.exists(path):
                return None, None

            ext = os.path.splitext(path)[1].lower()
            all_rows = []

            if ext == '.xls':
                import xlrd
                wb = xlrd.open_workbook(path)
                sheet = wb.sheet_by_index(0)
                for r in range(sheet.nrows):
                    all_rows.append(sheet.row_values(r))
            else:
                wb = openpyxl.load_workbook(path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(values_only=True):
                    all_rows.append(list(row))

            df_rows = []
            v_class_1_name = 'VL'
            v_class_2_name = 'PL'

            r = 0
            while r < len(all_rows):
                row_vals = all_rows[r]
                if not row_vals or len(row_vals) < 26:
                    r += 1
                    continue

                # Check if this row is a time header (e.g. 00:00, 01:00... or 0:00, 1:00... or 00h...)
                time_cell = str(row_vals[2]).strip() if len(row_vals) > 2 else ''
                if time_cell in ['00:00', '0:00', '00h', '0h', '00:00:00']:
                    # Next rows are Class 1, Class 2, TV
                    current_date = None
                    if r + 1 < len(all_rows):
                        r1 = all_rows[r + 1]
                        date_str = str(r1[0]).strip() if len(r1) > 0 and r1[0] is not None else ''
                        v1 = str(r1[1]).strip() if len(r1) > 1 and r1[1] is not None else 'VL'
                        if v1: v_class_1_name = v1

                        # Try parsing date
                        for fmt in ['%d/%m/%y', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d-%m-%y']:
                            try:
                                current_date = datetime.strptime(date_str, fmt)
                                break
                            except ValueError:
                                pass

                        if current_date:
                            for h in range(24):
                                raw_val = r1[2 + h] if 2 + h < len(r1) else 0
                                cnt = int(float(raw_val)) if raw_val not in ['', None] else 0
                                ts = current_date + timedelta(hours=h)
                                df_rows.append({
                                    'sensor_id': 1,
                                    'direction': 'Sens 1',
                                    'vehicle_class': v_class_1_name,
                                    'timestamp': ts,
                                    'count': cnt,
                                    **{f'Bin{i+1}': 0 for i in range(12)}
                                })

                    if r + 2 < len(all_rows) and current_date:
                        r2 = all_rows[r + 2]
                        v2 = str(r2[1]).strip() if len(r2) > 1 and r2[1] is not None else 'PL'
                        if v2: v_class_2_name = v2

                        for h in range(24):
                            raw_val = r2[2 + h] if 2 + h < len(r2) else 0
                            cnt = int(float(raw_val)) if raw_val not in ['', None] else 0
                            ts = current_date + timedelta(hours=h)
                            df_rows.append({
                                'sensor_id': 2,
                                'direction': 'Sens 1',
                                'vehicle_class': v_class_2_name,
                                'timestamp': ts,
                                'count': cnt,
                                **{f'Bin{i+1}': 0 for i in range(12)}
                            })

                    r += 4 # Move past Header, Class 1, Class 2, TV
                else:
                    r += 1

            if not df_rows:
                return None, None

            df = pd.DataFrame(df_rows)
            df['timestamp'] = pd.to_datetime(df['timestamp'])

            min_ts = df['timestamp'].min()
            max_ts = df['timestamp'].max()

            metadata = {
                'year': min_ts.year,
                'month': min_ts.month,
                'day': min_ts.day,
                'start_hour': min_ts.hour,
                'start_minute': min_ts.minute,
                'interval_minutes': 60,
                'num_sensors': 2,
                'mode': 3,
                'has_velocity': False,  # Count-only format
                'speed_bins': [30, 40, 50, 60, 70, 80, 90, 100, 110, 120, 130, 150],
                'speed_bin_centers': [15, 35, 45, 55, 65, 75, 85, 95, 105, 115, 125, 140],
                'vl_columns': list(range(12)),
                'pl_columns': list(range(12)),
                'gps_coordinates': None,
                'sensor_map': {
                    0: {'direction': 'Sens 1', 'class': v_class_1_name},
                    1: {'direction': 'Sens 1', 'class': v_class_2_name}
                },
                'start_datetime': min_ts.strftime('%d/%m/%Y %H:%M'),
                'end_datetime': max_ts.strftime('%d/%m/%Y %H:%M'),
                'raw_data': [],
                'rows_per_block': len(df) // 2 if len(df) >= 2 else len(df),
                'num_data_rows': len(df),
                'filtered_artificial_rows': 0
            }

            return df, metadata

        except Exception as e:
            print(f"Error parsing Raw Data file: {e}")
            return None, None

    @staticmethod
    def export_raw_data_weekly_matrix(df: pd.DataFrame, metadata: Dict[str, Any], filepath: str) -> bool:
        """
        Export traffic data to Excel in the Weekly Matrix format matching Raw-data-format.XLS.
        """
        try:
            if df is None or df.empty:
                return False

            wb = openpyxl.Workbook()
            ws = wb.active

            # Determine ISO week number
            min_date = df['timestamp'].min()
            iso_week = min_date.isocalendar()[1] if pd.notnull(min_date) else 1
            ws.title = f"Semaine_{iso_week:02d}"

            # Identify vehicle classes present
            classes = sorted(df['vehicle_class'].unique())
            v_class_1 = 'VL'
            v_class_2 = 'PL'
            if '2RM' in classes or '2R' in classes:
                v_class_1 = '2RM' if '2RM' in classes else classes[0]
                v_class_2 = '2R' if '2R' in classes else (classes[1] if len(classes) > 1 else 'PL')
            elif 'VL' in classes:
                v_class_1 = 'VL'
                v_class_2 = 'PL' if 'PL' in classes else (classes[1] if len(classes) > 1 else 'PL')

            # Get sorted unique dates
            dates = sorted(df['timestamp'].dt.date.unique())

            current_row = 1
            for d in dates:
                day_name = FRENCH_DAYS[d.weekday()]
                
                # Header row: [day_name, ' ', '00:00', '01:00', ..., '23:00', 'Total']
                header_cols = [day_name, ' '] + [f"{h:02d}:00" for h in range(24)] + ['Total']
                for col_idx, val in enumerate(header_cols, 1):
                    ws.cell(row=current_row, column=col_idx, value=val)

                # Fetch daily slice
                daily_df = df[df['timestamp'].dt.date == d]

                # Class 1 row
                c1_counts = []
                c1_df = daily_df[daily_df['vehicle_class'] == v_class_1]
                for h in range(24):
                    val = c1_df[c1_df['timestamp'].dt.hour == h]['count'].sum()
                    c1_counts.append(int(val))
                c1_total = sum(c1_counts)
                c1_row = [d.strftime('%d/%m/%y'), v_class_1] + c1_counts + [c1_total]
                for col_idx, val in enumerate(c1_row, 1):
                    ws.cell(row=current_row + 1, column=col_idx, value=val)

                # Class 2 row
                c2_counts = []
                c2_df = daily_df[daily_df['vehicle_class'] == v_class_2]
                for h in range(24):
                    val = c2_df[c2_df['timestamp'].dt.hour == h]['count'].sum()
                    c2_counts.append(int(val))
                c2_total = sum(c2_counts)
                c2_row = [' ', v_class_2] + c2_counts + [c2_total]
                for col_idx, val in enumerate(c2_row, 1):
                    ws.cell(row=current_row + 2, column=col_idx, value=val)

                # TV row (Class 1 + Class 2)
                tv_counts = [c1_counts[h] + c2_counts[h] for h in range(24)]
                tv_total = c1_total + c2_total
                tv_row = [' ', 'TV'] + tv_counts + [tv_total]
                for col_idx, val in enumerate(tv_row, 1):
                    ws.cell(row=current_row + 3, column=col_idx, value=val)

                # Advance past 4 rows + 1 blank row separator
                current_row += 5

            # Save workbook
            wb.save(filepath)
            return True

        except Exception as e:
            print(f"Error exporting Weekly Matrix Raw Data: {e}")
            return False
