"""
Comprehensive Verification Suite for DataFerns V3 Capabilities
Tests:
1. Raw Data Upload (XLS / XLSX weekly matrix)
2. Raw Data Download (Weekly Matrix format & Standard format)
3. Automatic Velocity Detection (With vs Without Velocity)
4. Adaptive Report Generation (Standard 4-sheet report vs Adapted 3-sheet report)
5. Backward Compatibility (CSV Ingress/Egress, FIM Parser, TMJ Analytics)
"""
import os
import sys
import shutil
import unittest
import pandas as pd
import openpyxl

# Add project root to sys.path so src is importable as a package
project_root = os.path.dirname(__file__)
sys.path.insert(0, project_root)

from src.services.raw_data_service import RawDataService
from src.services.excel_service import ExcelService
from src.services.csv_ingress_service import CsvIngressService
from src.services.report_service import ReportService
from src.core.fim_parser import parse_fim_file
from src.core.analytics import AnalyticsEngine

class TestDataFernsV3(unittest.TestCase):
    
    @classmethod
    def setUpClass(cls):
        cls.ref_dir = r"C:\Users\suppo\Downloads\DataFIMLoader-V3\DataFIMLoader-V3"
        cls.raw_xls_path = os.path.join(cls.ref_dir, "Raw-data-format.XLS")
        cls.scratch_dir = os.path.join(os.path.dirname(__file__), "scratch", "test_out")
        os.makedirs(cls.scratch_dir, exist_ok=True)

    def test_01_raw_data_ingress(self):
        """Test loading Raw-data-format.XLS into application DataFrame & metadata"""
        print("\n--- Running Test 01: Raw Data Ingress ---")
        df, meta = RawDataService.parse_raw_data_file(self.raw_xls_path)
        self.assertIsNotNone(df, "DataFrame should not be None")
        self.assertIsNotNone(meta, "Metadata should not be None")
        self.assertFalse(df.empty, "DataFrame should not be empty")
        
        # Verify columns
        required_cols = ['sensor_id', 'direction', 'vehicle_class', 'timestamp', 'count']
        for col in required_cols:
            self.assertIn(col, df.columns)
            
        # Verify metadata
        self.assertFalse(meta.get('has_velocity', True), "Raw data XLS should be flagged as has_velocity=False")
        self.assertEqual(meta.get('interval_minutes'), 60)
        self.assertIn('start_datetime', meta)
        self.assertIn('end_datetime', meta)
        print(f"[OK] Parsed {len(df)} records from {meta['start_datetime']} to {meta['end_datetime']}")

    def test_02_raw_data_weekly_export(self):
        """Test exporting DataFrame to Weekly Matrix format and validating structure"""
        print("\n--- Running Test 02: Raw Data Weekly Export ---")
        df, meta = RawDataService.parse_raw_data_file(self.raw_xls_path)
        
        export_path = os.path.join(self.scratch_dir, "Exported_Raw_Weekly.xlsx")
        success = RawDataService.export_raw_data_weekly_matrix(df, meta, export_path)
        self.assertTrue(success, "Export should return True")
        self.assertTrue(os.path.exists(export_path), "Exported file must exist")
        
        # Re-read exported file to verify integrity
        wb = openpyxl.load_workbook(export_path, data_only=True)
        sheet = wb.active
        self.assertTrue(sheet.title.startswith("Semaine_"), f"Sheet title should start with Semaine_, got {sheet.title}")
        
        # Check first header row
        header_vals = [sheet.cell(1, c).value for c in range(1, 28)]
        self.assertEqual(header_vals[0], 'mardi')
        self.assertEqual(header_vals[2], '00:00')
        self.assertEqual(header_vals[-1], 'Total')
        
        # Check VL row
        vl_row = [sheet.cell(2, c).value for c in range(1, 28)]
        self.assertEqual(vl_row[1], 'VL')
        self.assertEqual(vl_row[-1], 4928) # Day total from reference
        print(f"[OK] Exported weekly matrix matches expected layout and totals ({sheet.title})")

    def test_03_standard_raw_export(self):
        """Test standard multi-sheet Excel export backward compatibility"""
        print("\n--- Running Test 03: Standard Raw Excel Export ---")
        df, meta = RawDataService.parse_raw_data_file(self.raw_xls_path)
        
        path = ExcelService.export_raw_data(df, meta, self.scratch_dir)
        self.assertTrue(os.path.exists(path))
        
        wb = openpyxl.load_workbook(path, data_only=True)
        self.assertIn('Métadonnées', wb.sheetnames)
        self.assertIn('Comptages Sens 1', wb.sheetnames)
        print(f"[OK] Standard multi-sheet export generated: {wb.sheetnames}")

    def test_04_no_velocity_report_generation(self):
        """Test adaptive report generation for dataset without velocity"""
        print("\n--- Running Test 04: No-Velocity Adaptive Report Generation ---")
        df, meta = RawDataService.parse_raw_data_file(self.raw_xls_path)
        
        # Add Sens 2 for complete 3-sheet validation
        df_s2 = df.copy()
        df_s2['direction'] = 'Sens 2'
        df_full = pd.concat([df, df_s2], ignore_index=True)
        
        settings = {
            'site_name': 'Troyes',
            'voie': 'Av. Robert Schumann',
            'sect_info': 'C01',
            'sens': 'Vers Av. Jules Guesde',
            'sens2_name': 'De Av. Jules Guesde',
            'vmax': 50
        }
        
        report_path = ReportService.generate_report(df_full, meta, settings, self.scratch_dir)
        self.assertTrue(os.path.exists(report_path), "Report file must be created")
        
        wb = openpyxl.load_workbook(report_path, data_only=True)
        expected_sheets = ['Synthèse Sens1', 'Synthèse Sens2', 'Synthèse Sens3']
        self.assertEqual(wb.sheetnames, expected_sheets)
        
        # Validate Sheet 1 content
        ws = wb['Synthèse Sens1']
        self.assertEqual(ws['F1'].value, 'COMMUNE DE Troyes')
        self.assertEqual(ws['R1'].value, 'Av. Robert Schumann')
        self.assertEqual(ws['L2'].value, 'C01')
        self.assertEqual(ws['N2'].value, 1)
        self.assertEqual(ws['R2'].value, 'Vers Av. Jules Guesde')
        self.assertEqual(ws['F3'].value, '(DEBIT VL/PL)')
        
        # Check KPI values
        self.assertIsNotNone(ws['D55'].value, "TMJO TVC should be populated")
        self.assertIsNotNone(ws['H55'].value, "TMJA TVC should be populated")
        self.assertIsNotNone(ws['D59'].value, "TMJO PL should be populated")
        self.assertIsNotNone(ws['H59'].value, "TMJA PL should be populated")
        
        # Check charts presence
        self.assertGreaterEqual(len(ws._charts), 2, "Report must contain distribution charts on Page 2")
        print(f"[OK] Adapted 3-sheet report generated successfully with correct metadata, grids, and {len(ws._charts)} charts")

    def test_05_velocity_presence_and_standard_report(self):
        """Test FIM file with velocity generates standard 4-sheet report"""
        print("\n--- Running Test 05: Standard Report with Velocity ---")
        fim_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "test_suspicious_data.fim")
        if os.path.exists(fim_path):
            df, meta = parse_fim_file(fim_path)
            self.assertIsNotNone(df)
            self.assertTrue(meta.get('has_velocity', True), "FIM with speed bins must have has_velocity=True")
            
            # Slice to 7 days for fast report generation
            days = sorted(df['timestamp'].dt.date.unique())[:7]
            df_slice = df[df['timestamp'].dt.date.isin(days)].copy()
            
            settings = {
                'site_name': 'TestSite',
                'vmax': 50,
                'sect_info': 'Sect: 0170 / Ind: 11 / Count: 0000',
                'periods': ['07:00-09:00', '12:00-14:00', '17:00-19:00'],
                'sens': 'Sens 1'
            }
            
            report_path = ReportService.generate_report(df_slice, meta, settings, self.scratch_dir)
            self.assertTrue(os.path.exists(report_path))
            
            wb = openpyxl.load_workbook(report_path, data_only=True)
            self.assertIn('Synthese_des_donnees', wb.sheetnames)
            self.assertIn('Sens 1', wb.sheetnames)
            print(f"[OK] Standard full velocity report generated with sheets: {wb.sheetnames}")

    def test_06_csv_ingress_egress_compatibility(self):
        """Test CSV ingress and egress with automatic metadata reconstruction"""
        print("\n--- Running Test 06: CSV Ingress/Egress Compatibility ---")
        df, meta = RawDataService.parse_raw_data_file(self.raw_xls_path)
        
        csv_path = os.path.join(self.scratch_dir, "test_raw.csv")
        success = CsvIngressService.export_raw_to_csv(df, meta, csv_path)
        self.assertTrue(success)
        
        df_read, meta_read = CsvIngressService.parse_csv_file(csv_path)
        self.assertIsNotNone(df_read)
        self.assertEqual(len(df_read), len(df))
        self.assertFalse(meta_read.get('has_velocity', True))
        print(f"[OK] CSV Ingress/Egress cycle completed successfully ({len(df_read)} rows)")

if __name__ == '__main__':
    unittest.main()
